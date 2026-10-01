"""Excel product import with a review step and atomic store-scoped writes."""
import io
import logging
import time
import uuid
import zipfile
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.contrib import messages
from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from django.http import FileResponse, HttpResponse
from django.shortcuts import redirect, render

from .models import Category, Product, Store
from .views import can_manage_products, get_current_store, has_product_management_access

logger = logging.getLogger(__name__)
HEADERS = ["Name", "Category", "Cost Price", "Selling Price",
           "Stock", "Low Stock Threshold", "Barcode"]
STAGE_KEY = "bulk_product_preview"
MAX_ROWS = 1000


def _access(request, require_subscription=True):
    if not request.user.is_active or not can_manage_products(request.user):
        return None, HttpResponse("You do not have permission to import products.", status=403)
    store = get_current_store(request)
    if store is None:
        messages.error(request, "No active store is assigned to your account.")
        return None, redirect("dashboard")
    if require_subscription and not has_product_management_access(store):
        messages.warning(
            request, "Your trial or subscription has expired. Subscribe to import products.")
        return None, redirect("subscription")
    return store, None


def download_template(request):
    _, response = _access(request, require_subscription=False)
    if response is not None:
        return response
    path = Path(__file__).resolve().parent / "bulk_products_template.xlsx"
    if not path.is_file():
        messages.error(
            request, "The Excel template is missing. Ask your administrator to install it.")
        return redirect("add_product")
    return FileResponse(
        path.open("rb"), as_attachment=True, filename="Kounter_Product_Template.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _number(value, field, integer=False, default=None):
    if value is None or (isinstance(value, str) and not value.strip()):
        if default is not None:
            return default
        raise ValueError(f"{field} is required.")
    if isinstance(value, bool):
        raise ValueError(f"{field} must be a number.")
    try:
        number = Decimal(str(value).strip())
        maximum = Decimal("2147483647") if integer else Decimal(
            "9999999999.99")
        if not number.is_finite() or number < 0 or number > maximum:
            raise ValueError
        if integer:
            if number != number.to_integral_value():
                raise ValueError
            return int(number)
        if number != number.quantize(Decimal("0.01")):
            raise ValueError
        return str(number.quantize(Decimal("0.01")))
    except (InvalidOperation, ValueError, TypeError):
        description = "a non-negative whole number" if integer else "a non-negative price with at most 2 decimal places"
        raise ValueError(
            f"{field} must be {description} within the supported range.")


def validate_rows(raw_rows, store):
    """Normalize values and identify every invalid row without changing the database."""
    existing_names = {name.casefold() for name in Product.objects.filter(
        store=store).values_list("name", flat=True)}
    submitted_barcodes = [values[6].strip() for _, values in raw_rows
                          if isinstance(values[6], str) and values[6].strip()]
    existing_barcodes = set(Product.objects.filter(
        barcode__in=submitted_barcodes).values_list("barcode", flat=True))
    seen_names, seen_barcodes = set(), set()
    rows, errors = [], []
    for row_number, values in raw_rows:
        try:
            name, category, cost, selling, stock, threshold, barcode = values
            if not isinstance(name, str) or not name.strip():
                raise ValueError("Name is required and must be text.")
            name = name.strip()
            if len(name) > 200:
                raise ValueError("Name must not exceed 200 characters.")
            key = name.casefold()
            if key in existing_names or key in seen_names:
                raise ValueError(
                    "Product name already exists in your store or repeats in this file.")
            if category is None:
                category = ""
            if not isinstance(category, str) or len(category.strip()) > 100:
                raise ValueError(
                    "Category must be text of at most 100 characters.")
            if barcode is None:
                barcode = ""
            if not isinstance(barcode, str):
                raise ValueError(
                    "Barcode must be stored as Text in Excel so leading zeros are preserved.")
            barcode = barcode.strip()
            if len(barcode) > 100:
                raise ValueError("Barcode must not exceed 100 characters.")
            if barcode and (barcode in seen_barcodes or barcode in existing_barcodes):
                raise ValueError(
                    "Barcode already exists or repeats in this file.")
            row = {
                "row_number": row_number, "name": name, "category": category.strip(),
                "cost_price": _number(cost, "Cost Price"),
                "selling_price": _number(selling, "Selling Price"),
                "stock": _number(stock, "Stock", integer=True, default=0),
                "low_stock_threshold": _number(threshold, "Low Stock Threshold", integer=True, default=20),
                "barcode": barcode,
            }
            # Apply the model's own field validators, including database integer limits.
            product = Product(
                store=store, name=row["name"], cost_price=Decimal(
                    row["cost_price"]),
                selling_price=Decimal(row["selling_price"]), stock=row["stock"],
                low_stock_threshold=row["low_stock_threshold"], barcode=barcode or None,
                active=True,
            )
            product.full_clean(
                exclude=["category"], validate_unique=False, validate_constraints=False)
            seen_names.add(key)
            if barcode:
                seen_barcodes.add(barcode)
            rows.append(row)
        except (ValueError, ValidationError) as error:
            text = " ".join(error.messages) if isinstance(
                error, ValidationError) else str(error)
            errors.append(f"Row {row_number}: {text}")
    return rows, errors


def _read_file(upload):
    if not upload.name.lower().endswith(".xlsx"):
        raise ValueError(
            "Upload a modern Excel .xlsx file. Save older .xls files as .xlsx first.")
    if upload.size > 5 * 1024 * 1024:
        raise ValueError("The file must be no larger than 5 MB.")
    data = upload.read()
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 200 or sum(entry.file_size for entry in entries) > 20 * 1024 * 1024:
                raise ValueError(
                    "The workbook is too large. Use the supplied template.")
    except zipfile.BadZipFile:
        raise ValueError("This is not a readable .xlsx workbook.")

    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(
        data), read_only=True, data_only=False, keep_links=False)
    try:
        if "Products" not in workbook.sheetnames:
            raise ValueError(
                "The workbook needs a Products sheet. Use the supplied template.")
        sheet = workbook["Products"]
        # Stream actual XML rows instead of trusting the workbook's declared dimensions.
        sheet.reset_dimensions()
        iterator = sheet.iter_rows(max_col=8)
        first = next(iterator, ())
        headers = [str(cell.value or "").strip() for cell in first[:7]]
        if headers != HEADERS or (len(first) > 7 and first[7].value is not None):
            raise ValueError(
                "The header row has changed. Use the template's seven column headings.")
        raw_rows, errors = [], []
        for row_number, cells in enumerate(iterator, start=2):
            if row_number > MAX_ROWS + 1:
                raise ValueError(
                    "Use no more than 1,000 product rows per workbook.")
            values = [cell.value for cell in cells[:7]]
            if all(value is None or value == "" for value in values) and not cells[7].value:
                continue
            if cells[7].value is not None:
                errors.append(
                    f"Row {row_number}: use only the template's seven columns.")
            elif any(cell.data_type == "f" for cell in cells[:7]):
                errors.append(
                    f"Row {row_number}: enter values, not Excel formulas.")
            else:
                raw_rows.append((row_number, values))
        return raw_rows, errors
    finally:
        workbook.close()


def upload_products(request):
    store, response = _access(request)
    if response is not None:
        return response
    request.session.pop(STAGE_KEY, None)
    upload = request.FILES.get("product_file")
    if upload is None:
        messages.error(request, "Choose your filled Excel template first.")
        return redirect("add_product")
    try:
        raw_rows, errors = _read_file(upload)
        rows, row_errors = validate_rows(raw_rows, store)
        errors.extend(row_errors)
        if not rows and not errors:
            errors.append(
                "The Products sheet is empty. Fill it from row 2 and upload again.")
    except ValueError as error:
        rows, errors = [], [str(error)]
    except Exception:
        logger.exception(
            "Cannot read bulk product workbook for store %s", store.pk)
        rows, errors = [], ["The workbook could not be read. Save it as .xlsx and upload again."]
    token = ""
    if not errors:
        token = uuid.uuid4().hex
        request.session[STAGE_KEY] = {
            "owner": request.user.pk, "store": store.pk,
            "created": time.time(), "token": token, "rows": rows,
        }
    return render(request, "core/bulk_product_preview.html", {
        "store": store, "rows": rows, "errors": errors, "token": token,
    })


def confirm_products(request):
    store, response = _access(request)
    if response is not None:
        return response
    stage = request.session.get(STAGE_KEY)
    if not (
        stage and stage.get("owner") == request.user.pk
        and stage.get("store") == store.pk
        and stage.get("token") == request.POST.get("token")
        and 0 <= time.time() - stage.get("created", 0) <= 15 * 60
    ):
        request.session.pop(STAGE_KEY, None)
        messages.error(
            request, "The import preview has expired or is invalid. Upload the file again.")
        return redirect("add_product")

    rows = stage["rows"]
    try:
        with transaction.atomic():
            # Serializes imports into this store and protects repeat submissions.
            locked_store = Store.objects.select_for_update().get(pk=store.pk)
            if not locked_store.active or not has_product_management_access(locked_store):
                raise ValueError(
                    "Your store or subscription is no longer active.")
            raw = [(row["row_number"], [row["name"], row["category"], row["cost_price"],
                    row["selling_price"], row["stock"], row["low_stock_threshold"], row["barcode"]]) for row in rows]
            normalized, errors = validate_rows(raw, locked_store)
            if errors:
                raise ValueError(" ".join(errors))
            category_map = {category.name.casefold(
            ): category for category in Category.objects.filter(store=locked_store)}
            missing_categories = {}
            for row in normalized:
                key = row["category"].casefold()
                if key and key not in category_map:
                    missing_categories.setdefault(key, row["category"])
            if missing_categories:
                Category.objects.bulk_create([
                    Category(store=locked_store, name=name) for name in missing_categories.values()
                ], batch_size=200)
                category_map = {category.name.casefold(
                ): category for category in Category.objects.filter(store=locked_store)}
            products = [Product(
                store=locked_store, category=category_map.get(
                    row["category"].casefold()),
                name=row["name"], cost_price=Decimal(row["cost_price"]),
                selling_price=Decimal(row["selling_price"]), stock=row["stock"],
                low_stock_threshold=row["low_stock_threshold"], barcode=row["barcode"] or None,
                active=True,
            ) for row in normalized]
            Product.objects.bulk_create(products, batch_size=200)
    except (ValueError, IntegrityError) as error:
        request.session.pop(STAGE_KEY, None)
        if isinstance(error, IntegrityError):
            logger.exception("Bulk import rolled back for store %s", store.pk)
            text = "A product or barcode conflicted with another update. No products were imported. Upload again."
        else:
            text = str(error)
        return render(request, "core/bulk_product_preview.html", {
            "store": store, "rows": rows, "errors": [text], "token": "",
        })
    request.session.pop(STAGE_KEY, None)
    messages.success(
        request, f"{len(rows)} products imported successfully into {store.name}.")
    return redirect("active_products")
